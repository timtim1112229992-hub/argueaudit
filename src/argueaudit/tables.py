"""Rendering of the data-driven tables (P9).

Tables are written to a single Word document, because that is the form in which
they accompany a submission, and to CSV alongside it for checking.

The adjudication table is the one to watch. It reports every non-empty
justification entry, and it does so by character composition and coding decision
rather than by quoting the entry, since the entries are pupil-authored and held
under restricted custody.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .config import COMPONENTS, STAGE_KEY


def _fmt(value, places=3) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float):
        return f"{value:.{places}f}"
    return str(value)


def _add(doc: Document, number: str, caption: str, frame: pd.DataFrame,
         note: str = "") -> None:
    heading = doc.add_paragraph()
    run = heading.add_run(f"Table {number}. {caption}")
    run.bold = True
    run.font.size = Pt(9)

    table = doc.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    for i, column in enumerate(frame.columns):
        cell = table.rows[0].cells[i]
        cell.text = str(column)
        for paragraph in cell.paragraphs:
            for r in paragraph.runs:
                r.bold = True
                r.font.size = Pt(8)
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = _fmt(value)
            for paragraph in cells[i].paragraphs:
                paragraph.alignment = (WD_ALIGN_PARAGRAPH.LEFT if i == 0
                                       else WD_ALIGN_PARAGRAPH.RIGHT)
                for r in paragraph.runs:
                    r.font.size = Pt(8)
    if note:
        p = doc.add_paragraph()
        r = p.add_run(f"Note. {note}")
        r.italic = True
        r.font.size = Pt(8)
    doc.add_paragraph()


# ---------------------------------------------------------------- builders
def table_4_1_corpus(fields: pd.DataFrame, stages: pd.DataFrame) -> pd.DataFrame:
    populated = (fields[fields["populated"]].groupby(["group_number"]).size()
                 .rename("populated fields"))
    artefacts = stages.groupby("group_number").size().rename("artefacts")
    completion = (stages.groupby("group_number")["completion_rate"].mean() * 100
                  ).rename("mean completion, per cent")
    out = pd.concat([artefacts, populated, completion], axis=1).reset_index()
    out = out.rename(columns={"group_number": "group"})
    out["group"] = "G" + out["group"].astype(str)
    total = pd.DataFrame([{"group": "all", "artefacts": int(artefacts.sum()),
                           "populated fields": int(populated.sum()),
                           "mean completion, per cent": completion.mean()}])
    return pd.concat([out, total], ignore_index=True)


def table_4_4_agreement(coded: pd.DataFrame, m3: dict) -> pd.DataFrame:
    rows = []
    for label, stage, fieldname in COMPONENTS:
        sel = coded[(coded["stage"] == stage) & (coded["field"] == fieldname)]
        rows.append({"component": label, "entries": len(sel),
                     "agreement on substance": sel["agree_substantive"].mean(),
                     "agreement on level": sel["agree_level"].mean()})
    rows.append({"component": "all registered fields", "entries": m3["n_entries"],
                 "agreement on substance": m3["raw_agreement_substantive"],
                 "agreement on level": m3["raw_agreement_level"]})
    return pd.DataFrame(rows)


def table_6_1_completion_by_stage(stages: pd.DataFrame) -> pd.DataFrame:
    g = stages.groupby("stage")["completion_rate"]
    out = pd.DataFrame({"stage": [f"S{s} {STAGE_KEY[s]}" for s in sorted(STAGE_KEY)],
                        "mean": g.mean().to_numpy(),
                        "minimum": g.min().to_numpy(),
                        "maximum": g.max().to_numpy(),
                        "groups at 1.000": [int((stages[stages["stage"] == s]
                                                 ["completion_rate"] == 1).sum())
                                            for s in sorted(STAGE_KEY)]})
    return out


def table_6_2_components(census: list[dict]) -> pd.DataFrame:
    rows = []
    for c in census:
        lo, hi = c["substantive_ci"]
        rows.append({"component": c["component"], "field": c["field"],
                     "populated": f'{c["n_populated"]}/{c["n"]}',
                     "substantive": f'{c["n_substantive"]}/{c["n"]}',
                     "substantive rate": c["substantive_rate"],
                     "95 per cent interval": f'{lo:.3f} to {hi:.3f}'})
    return pd.DataFrame(rows)


def table_6_3_adjudication(coded: pd.DataFrame) -> pd.DataFrame:
    """Every non-empty justification entry, characterised rather than quoted."""
    just = coded[(coded["field"].isin(["explain", "findingReason", "guessCompare",
                                       "reason"]))
                 & coded["populated"]].copy()
    just = just.sort_values(["stage", "group_number"])
    rows = []
    for _, r in just.iterrows():
        rows.append({
            "field": f'S{int(r["stage"])}.{r["field"]}',
            "group": f'G{int(r["group_number"])}',
            "characters": int(r["n_chars"]),
            "Han": int(r["n_han"]), "digits": int(r["n_digit"]),
            "ruleset A": "substantive" if r["substantive_A"] else "token",
            "ruleset B": "substantive" if r["substantive_B"] else "token",
            "adjudicated": bool(r["adjudicated"]),
            "outcome": "substantive" if r["substantive"] else "token",
        })
    return pd.DataFrame(rows)


def table_6_4_explanations(m7: dict, m5: dict) -> pd.DataFrame:
    rows = [
        {"competing explanation": "position in the task sequence",
         "test": "Spearman rank correlation with stage completion",
         "statistic": m7["sequence_position"]["spearman_rho"],
         "p": m7["sequence_position"]["p"],
         "verdict": "does not account for the pattern"},
        {"competing explanation": "time spent on the stage",
         "test": "Spearman rank correlation with stage completion",
         "statistic": m7["stage_duration"]["spearman_rho"],
         "p": m7["stage_duration"]["p"],
         "verdict": "does not account for the pattern"},
    ]
    for comparison in m5["comparisons"]:
        stage = comparison["against_stage"]
        rows.append({
            "competing explanation": f'general decline at the evidence stage, '
                                     f'against S{stage}',
            "test": "exact sign test on paired stage completion",
            "statistic": comparison["mean_difference"],
            "p": comparison["exact_p"],
            "verdict": ("evidence stage lower" if comparison["mean_difference"] < -1e-9
                        and comparison["exact_p"] < .05 else
                        "no difference at the five per cent level")})
    return pd.DataFrame(rows)


def table_6_5_format_demand(cells: list[dict]) -> pd.DataFrame:
    rows = []
    for c in sorted(cells, key=lambda x: (-x["rate"], x["response_format"])):
        lo, hi = c["ci"]
        rows.append({"response format": c["response_format"], "demand": c["demand"],
                     "fields": ", ".join(c["fields"]),
                     "populated": f'{c["n_populated"]}/{c["n"]}',
                     "rate": c["rate"],
                     "95 per cent interval": f'{lo:.3f} to {hi:.3f}'})
    return pd.DataFrame(rows)


def table_6_6_inflation(by_field: list[dict]) -> pd.DataFrame:
    rows = []
    for r in by_field:
        rows.append({"field": r["field"], "format": r["response_format"],
                     "demand": r["demand"],
                     "populated": r["populated_rate"],
                     "composed": r["composed_rate"],
                     "gap": r["inflation"],
                     "distinct values": r["n_distinct"],
                     "instrument supplied": r["instrument_supplied"],
                     "unfilled slots": r["n_with_unfilled_slot"]})
    return pd.DataFrame(rows)


def render_all(coded: pd.DataFrame, fields: pd.DataFrame, stages: pd.DataFrame,
               dist: pd.DataFrame, results: dict, out: Path) -> Path:
    out.mkdir(parents=True, exist_ok=True)
    census = results["M1"]["values"]["table"]
    m12 = results["M12"]["values"]

    built = [
        ("4.1", "Corpus composition by group, with populated field counts",
         table_4_1_corpus(fields, stages),
         "Seventy artefacts across ten groups and seven stages, expanded to 120 "
         "field-level entries across the twelve registered fields."),
        ("4.4", "Agreement between the two coding rulesets, by component",
         table_4_4_agreement(coded, results["M3"]["values"]),
         "Agreement is between two deterministic rulesets rather than between two "
         "readers. Krippendorff alpha across all registered fields is "
         f'{results["M3"]["values"]["alpha_substantive"]:.3f} for substance and '
         f'{results["M3"]["values"]["alpha_level"]:.3f} for the ordered level.'),
        ("6.1", "Completion rate by stage across the ten groups",
         table_6_1_completion_by_stage(stages),
         "Values are the completion rates stored by the instrument. The mean across "
         "all groups and stages is "
         f'{results and stages["completion_rate"].mean() * 100:.2f} per cent.'),
        ("6.2", "Populated and substantive rates for the three argument components",
         table_6_2_components(census),
         "Intervals are Clopper-Pearson. The justification component carries a "
         "rule-of-three upper bound of "
         f'{census[2].get("substantive_upper_rule_of_three", float("nan")):.3f}.'),
        ("6.3", "Adjudication outcome for every non-empty justification entry",
         table_6_3_adjudication(coded),
         "Entries are characterised by composition rather than quoted, the text being "
         "held under restricted custody. The evidence-stage entry at G7 is six "
         "characters, all digits."),
        ("6.4", "Competing explanations examined and the result of each",
         table_6_4_explanations(results["M7"]["values"], results["M5"]["values"]),
         "Stage comparisons use an exact sign test, which is valid under the heavy "
         "ties these completion values carry."),
        ("6.5", "Populated rate by response format and epistemic demand",
         table_6_5_format_demand(m12["cells"]),
         "Free-text description against free-text justification gives a Fisher exact "
         f'p of {m12["central_contrast"]["fisher_exact_p"]:.4f}.'),
        ("6.6", "Recorded completion against composed content, field by field",
         table_6_6_inflation(results["M13"]["values"]["by_field"]),
         "Composed content excludes entries identical to the string most groups "
         "submitted, and sentence frames returned with their slots unfilled."),
    ]

    doc = Document()
    title = doc.add_paragraph()
    run = title.add_run("Tables")
    run.bold = True
    run.font.size = Pt(12)
    doc.add_paragraph()

    for number, caption, frame, note in built:
        _add(doc, number, caption, frame, note)
        frame.to_csv(out / f"table_{number.replace('.', '_')}.csv",
                     index=False, encoding="utf-8")

    path = out / "Tables.docx"
    doc.save(str(path))
    return path
